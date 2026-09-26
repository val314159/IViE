#!/usr/bin/env python3

import csv
import io
import os
import re
import shutil
import sys
import zipfile

from collections import OrderedDict
from datetime import date, datetime
from pathlib import Path, PurePosixPath

from openpyxl import load_workbook


ROOT = Path("data/ercot")

META_COLUMNS = [
    "_publication_date",
    "_publication_time",
    "_source_zip",
    "_source_member",
    "_extracted2_file",
    "_sheet",
]

# Handles:
#   .20260430.160053.
#   _20260430_160053_
#   -20260430-160053-
REPORT_TS_RE = re.compile(
    r"(?<!\d)(20\d{6})[._-](\d{6})(?!\d)"
)

REPORT_DATE_RE = re.compile(
    r"(?<!\d)(20\d{6})(?!\d)"
)

OUTER_MONTH_RE = re.compile(
    r"(20\d{2}-\d{2})"
)


# ----------------------------------------------------------------------
# Month handling
# ----------------------------------------------------------------------

def parse_month(s):
    y, m = map(int, s.split("-"))
    return y, m


def month_string(y, m):
    return f"{y:04d}-{m:02d}"


def shift_month(y, m, delta):
    n = y * 12 + (m - 1) + delta
    return n // 12, n % 12 + 1


def month_range(start, end):
    y, m = start

    while (y, m) <= end:
        yield y, m
        y, m = shift_month(y, m, 1)


def find_latest_complete_month(root):
    found = []

    for p in root.iterdir():
        if not p.is_dir():
            continue

        if not re.fullmatch(r"\d{4}-\d{2}", p.name):
            continue

        needed = [
            p / "prices.csv",
            p / "demand.csv",
            p / "fuel_mix.csv",
        ]

        if all(x.exists() for x in needed):
            found.append(parse_month(p.name))

    if not found:
        raise RuntimeError(
            "No YYYY-MM directories containing "
            "prices.csv, demand.csv, and fuel_mix.csv found"
        )

    return max(found)


# ----------------------------------------------------------------------
# Timestamp detection
# ----------------------------------------------------------------------

def timestamp_from_name(name):
    """
    Try to recover ERCOT publication date/time from a filename.
    """

    m = REPORT_TS_RE.search(name)

    if m:
        ds, ts = m.groups()

        return (
            f"{ds[0:4]}-{ds[4:6]}-{ds[6:8]}",
            f"{ts[0:2]}:{ts[2:4]}:{ts[4:6]}",
        )

    # Some old files may have a date but no publication time.
    m = REPORT_DATE_RE.search(name)

    if m:
        ds = m.group(1)

        return (
            f"{ds[0:4]}-{ds[4:6]}-{ds[6:8]}",
            "",
        )

    return None


def timestamp_for_record(record):
    """
    Prefer the inner ZIP filename.

    If that fails, try the actual extracted member filenames.
    This gives us a chance of handling older naming conventions.
    """

    ts = timestamp_from_name(record["zip_path"].name)

    if ts:
        return ts

    for member in record["members"]:
        ts = timestamp_from_name(member.name)

        if ts:
            return ts

    return None


def month_from_timestamp(ts):
    if not ts:
        return None

    y, m, _ = map(int, ts[0].split("-"))

    return y, m


# ----------------------------------------------------------------------
# Generic helpers
# ----------------------------------------------------------------------

def value_to_string(v):
    if v is None:
        return ""

    if isinstance(v, (datetime, date)):
        return v.isoformat()

    return str(v).strip()


def unique_headers(values):
    used = {}
    result = []

    for i, value in enumerate(values, start=1):

        name = value_to_string(value)

        if not name:
            name = f"column_{i}"

        name = re.sub(r"\s+", " ", name)

        base = name

        if base in used:
            used[base] += 1
            name = f"{base}_{used[base]}"
        else:
            used[base] = 1

        result.append(name)

    return result


def add_column(columns, name):
    if name not in columns:
        columns[name] = None


def decode_csv(data):
    for encoding in (
        "utf-8-sig",
        "utf-8",
        "cp1252",
        "latin1",
    ):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            pass

    raise UnicodeDecodeError(
        "unknown",
        data,
        0,
        1,
        "Could not decode CSV",
    )


def write_csv_atomic(filename, rows, columns):
    filename.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    tmp = filename.with_name(
        f".{filename.name}.tmp.{os.getpid()}"
    )

    try:

        with tmp.open(
            "w",
            newline="",
            encoding="utf-8",
        ) as f:

            writer = csv.DictWriter(
                f,
                fieldnames=list(columns),
                extrasaction="ignore",
            )

            writer.writeheader()

            for row in rows:
                writer.writerow(row)

        os.replace(tmp, filename)

    finally:

        if tmp.exists():
            tmp.unlink()


# ----------------------------------------------------------------------
# Second-stage ZIP extraction
# ----------------------------------------------------------------------

def outer_month_for_zip(zip_path, extracted_root):
    """
    Usually:

      extracted/
          NP3-763-CD_2026-04_<hash>/
              inner.zip

    Recover 2026-04 from that outer directory.
    """

    try:
        rel = zip_path.relative_to(extracted_root)
    except ValueError:
        return None

    for part in rel.parts[:-1]:
        m = OUTER_MONTH_RE.search(part)

        if m:
            try:
                return parse_month(m.group(1))
            except Exception:
                pass

    return None


def safe_member_path(name):
    """
    Convert a ZIP member path into a safe relative path.
    Refuse ../ tricks or absolute paths.
    """

    p = PurePosixPath(name)

    if p.is_absolute():
        return None

    if ".." in p.parts:
        return None

    clean_parts = [
        part
        for part in p.parts
        if part not in ("", ".")
    ]

    if not clean_parts:
        return None

    return Path(*clean_parts)


def extract_one_inner_zip(
    zip_path,
    extracted_root,
    extracted2_root,
):
    """
    Example:

      extracted/
        OUTER/
          report.zip

    becomes:

      extracted2/
        OUTER/
          report/
            actual.csv

    or:

      extracted2/
        OUTER/
          report/
            actual.xlsx
    """

    rel = zip_path.relative_to(extracted_root)

    destination = (
        extracted2_root
        / rel.parent
        / zip_path.stem
    )

    destination.mkdir(
        parents=True,
        exist_ok=True,
    )

    members_written = []

    try:

        with zipfile.ZipFile(zip_path) as z:

            for info in z.infolist():

                if info.is_dir():
                    continue

                rel_member = safe_member_path(
                    info.filename
                )

                if rel_member is None:
                    print(
                        f"WARNING: unsafe ZIP member skipped: "
                        f"{zip_path}: {info.filename}"
                    )
                    continue

                target = destination / rel_member

                target.parent.mkdir(
                    parents=True,
                    exist_ok=True,
                )

                with z.open(info) as src:
                    with target.open("wb") as dst:
                        shutil.copyfileobj(src, dst)

                members_written.append(target)

    except zipfile.BadZipFile as e:

        print(
            f"WARNING: bad inner ZIP skipped: "
            f"{zip_path}: {e}"
        )

        return None

    return {
        "zip_path": zip_path,
        "destination": destination,
        "members": members_written,
    }


def extract_second_stage(
    dataset_name,
    start,
    end,
):
    """
    Extract inner ZIPs into extracted2.

    We include one outer-month on either side because ERCOT outer
    archive bundles overlap calendar boundaries.

    ZIPs whose outer month cannot be recognized are also extracted
    so they remain available for forensic inspection.
    """

    dataset_root = ROOT / dataset_name

    extracted_root = dataset_root / "extracted"
    extracted2_root = dataset_root / "extracted2"

    if not extracted_root.exists():
        print(
            f"WARNING: missing source directory: "
            f"{extracted_root}"
        )
        return []

    extracted2_root.mkdir(
        parents=True,
        exist_ok=True,
    )

    padded_start = shift_month(
        *start,
        -1,
    )

    padded_end = shift_month(
        *end,
        1,
    )

    all_zips = sorted(
        extracted_root.rglob("*.zip")
    )

    selected = []

    skipped_outside_range = 0
    unknown_outer_month = 0

    for zip_path in all_zips:

        outer_month = outer_month_for_zip(
            zip_path,
            extracted_root,
        )

        if outer_month is not None:

            if not (
                padded_start
                <= outer_month
                <= padded_end
            ):
                skipped_outside_range += 1
                continue

        else:

            # Extract unknown formats too.
            # They're exactly the files we may need to inspect.
            unknown_outer_month += 1

        selected.append(zip_path)

    print(
        f"{dataset_name}: "
        f"{len(selected)} inner ZIPs selected"
    )

    if skipped_outside_range:
        print(
            f"  skipped outside padded range: "
            f"{skipped_outside_range}"
        )

    if unknown_outer_month:
        print(
            f"  unknown outer-month naming, "
            f"still extracting: "
            f"{unknown_outer_month}"
        )

    records = []

    for i, zip_path in enumerate(
        selected,
        start=1,
    ):

        record = extract_one_inner_zip(
            zip_path,
            extracted_root,
            extracted2_root,
        )

        if record:
            records.append(record)

        if (
            i % 1000 == 0
            or i == len(selected)
        ):
            print(
                f"  extracted "
                f"{i}/{len(selected)}"
            )

    return records


# ----------------------------------------------------------------------
# Capacity parsing
# ----------------------------------------------------------------------

def read_capacity_record(record):
    ts = timestamp_for_record(record)

    if not ts:
        return [], [], None

    pub_date, pub_time = ts

    rows = []
    discovered_columns = []

    csv_members = [
        p
        for p in record["members"]
        if p.suffix.lower() == ".csv"
    ]

    for path in csv_members:

        try:
            raw = path.read_bytes()
            text = decode_csv(raw)

        except Exception as e:
            print(
                f"WARNING: cannot read capacity CSV "
                f"{path}: {e}"
            )
            continue

        reader = csv.DictReader(
            io.StringIO(text)
        )

        if not reader.fieldnames:
            continue

        source_columns = [
            str(x).strip()
            for x in reader.fieldnames
            if x is not None
        ]

        for c in source_columns:
            if c not in discovered_columns:
                discovered_columns.append(c)

        for source_row in reader:

            if not any(
                value_to_string(v)
                for v in source_row.values()
            ):
                continue

            row = {
                "_publication_date": pub_date,
                "_publication_time": pub_time,
                "_source_zip": str(
                    record["zip_path"]
                ),
                "_source_member": path.name,
                "_extracted2_file": str(path),
                "_sheet": "",
            }

            for key, value in source_row.items():

                if key is None:
                    continue

                row[
                    str(key).strip()
                ] = value_to_string(value)

            rows.append(row)

    return (
        rows,
        discovered_columns,
        ts,
    )


# ----------------------------------------------------------------------
# Outage parsing
# ----------------------------------------------------------------------

HEADER_HINTS = (
    "resource",
    "outage",
    "mw",
    "fuel",
    "start",
    "end",
    "effective",
    "available",
    "capacity",
)


def choose_excel_header(rows):
    best_index = None
    best_score = -1

    for i, row in enumerate(rows[:30]):

        values = [
            value_to_string(v)
            for v in row
        ]

        nonempty = sum(
            bool(v)
            for v in values
        )

        if nonempty < 2:
            continue

        text = " ".join(
            values
        ).lower()

        hints = sum(
            1
            for hint in HEADER_HINTS
            if hint in text
        )

        score = (
            hints * 100
            + nonempty
        )

        if score > best_score:
            best_score = score
            best_index = i

    return best_index


def read_outage_record(record):
    ts = timestamp_for_record(record)

    if not ts:
        return [], [], None

    pub_date, pub_time = ts

    rows = []
    discovered_columns = []

    xlsx_members = [
        p
        for p in record["members"]
        if p.suffix.lower() == ".xlsx"
    ]

    for path in xlsx_members:

        try:

            wb = load_workbook(
                path,
                read_only=True,
                data_only=True,
            )

        except Exception as e:

            print(
                f"WARNING: cannot open outage XLSX "
                f"{path}: {e}"
            )

            continue

        try:

            for ws in wb.worksheets:

                worksheet_rows = list(
                    ws.iter_rows(
                        values_only=True
                    )
                )

                if not worksheet_rows:
                    continue

                header_index = choose_excel_header(
                    worksheet_rows
                )

                if header_index is None:

                    print(
                        f"WARNING: no table header found: "
                        f"{path}:{ws.title}"
                    )

                    continue

                headers = unique_headers(
                    worksheet_rows[
                        header_index
                    ]
                )

                for c in headers:
                    if c not in discovered_columns:
                        discovered_columns.append(c)

                for values in worksheet_rows[
                    header_index + 1:
                ]:

                    values = list(values)

                    if not any(
                        value_to_string(v)
                        for v in values
                    ):
                        continue

                    if len(values) > len(headers):

                        for n in range(
                            len(headers),
                            len(values),
                        ):

                            h = (
                                f"column_{n + 1}"
                            )

                            headers.append(h)

                            if (
                                h
                                not in discovered_columns
                            ):
                                discovered_columns.append(
                                    h
                                )

                    row = {
                        "_publication_date": pub_date,
                        "_publication_time": pub_time,
                        "_source_zip": str(
                            record["zip_path"]
                        ),
                        "_source_member": path.name,
                        "_extracted2_file": str(path),
                        "_sheet": ws.title,
                    }

                    for i, value in enumerate(
                        values
                    ):

                        if i >= len(headers):
                            break

                        row[
                            headers[i]
                        ] = value_to_string(
                            value
                        )

                    rows.append(row)

        finally:
            wb.close()

    return (
        rows,
        discovered_columns,
        ts,
    )


# ----------------------------------------------------------------------
# Group reports by actual embedded report month
# ----------------------------------------------------------------------

def group_records(records, start, end):
    grouped = {}
    unknown = []

    seen_zip_names = set()
    duplicate_names = 0

    for record in records:

        # Same report may theoretically appear in overlapping
        # ERCOT archive bundles.
        #
        # Inner ZIP filename contains a unique report ID, so use
        # it as the de-duplication key.
        key = record["zip_path"].name

        if key in seen_zip_names:
            duplicate_names += 1
            continue

        seen_zip_names.add(key)

        ts = timestamp_for_record(record)

        if not ts:
            unknown.append(record)
            continue

        month = month_from_timestamp(ts)

        if (
            month is not None
            and start <= month <= end
        ):
            grouped.setdefault(
                month,
                [],
            ).append(record)

    for records_for_month in grouped.values():
        records_for_month.sort(
            key=lambda r:
                r["zip_path"].name
        )

    return (
        grouped,
        unknown,
        duplicate_names,
    )


# ----------------------------------------------------------------------
# Build monthly CSV files
# ----------------------------------------------------------------------

def build_capacity_month(
    month,
    records,
):
    y, m = month

    output = (
        ROOT
        / month_string(y, m)
        / "capacity.csv"
    )

    if not records:

        print(
            f"  capacity: NO SOURCE REPORTS "
            f"-- no CSV written"
        )

        return

    rows = []
    columns = OrderedDict()

    for c in META_COLUMNS:
        add_column(columns, c)

    parsed_reports = 0

    for record in records:

        (
            new_rows,
            source_columns,
            ts,
        ) = read_capacity_record(record)

        if ts is None:
            continue

        parsed_reports += 1

        for c in source_columns:
            add_column(columns, c)

        rows.extend(new_rows)

    write_csv_atomic(
        output,
        rows,
        columns,
    )

    print(
        f"  capacity: "
        f"{parsed_reports:5d} reports  "
        f"{len(rows):8d} rows -> "
        f"{output}"
    )


def build_outage_month(
    month,
    records,
):
    y, m = month

    output = (
        ROOT
        / month_string(y, m)
        / "outages.csv"
    )

    if not records:

        print(
            f"  outages:  NO SOURCE REPORTS "
            f"-- no CSV written"
        )

        return

    rows = []
    columns = OrderedDict()

    for c in META_COLUMNS:
        add_column(columns, c)

    parsed_reports = 0

    for record in records:

        (
            new_rows,
            source_columns,
            ts,
        ) = read_outage_record(record)

        if ts is None:
            continue

        parsed_reports += 1

        for c in source_columns:
            add_column(columns, c)

        rows.extend(new_rows)

    write_csv_atomic(
        output,
        rows,
        columns,
    )

    print(
        f"  outages:  "
        f"{parsed_reports:5d} reports  "
        f"{len(rows):8d} rows -> "
        f"{output}"
    )


# ----------------------------------------------------------------------
# Main
# ----------------------------------------------------------------------

def main():
    if not ROOT.exists():
        sys.exit(
            f"Missing ERCOT root: {ROOT}"
        )

    latest = find_latest_complete_month(
        ROOT
    )

    start = shift_month(
        *latest,
        -24,
    )

    print(
        f"Processing 25 months: "
        f"{month_string(*start)} through "
        f"{month_string(*latest)}"
    )

    print()
    print(
        "SECOND-STAGE EXTRACTION"
    )

    capacity_records = extract_second_stage(
        "capacity",
        start,
        latest,
    )

    outage_records = extract_second_stage(
        "outages",
        start,
        latest,
    )

    print()
    print(
        "GROUPING EXTRACTED REPORTS"
    )

    (
        capacity_by_month,
        capacity_unknown,
        capacity_duplicates,
    ) = group_records(
        capacity_records,
        start,
        latest,
    )

    (
        outage_by_month,
        outage_unknown,
        outage_duplicates,
    ) = group_records(
        outage_records,
        start,
        latest,
    )

    print(
        f"capacity reports with unrecognized date: "
        f"{len(capacity_unknown)}"
    )

    print(
        f"outage reports with unrecognized date: "
        f"{len(outage_unknown)}"
    )

    print(
        f"capacity duplicate report names skipped: "
        f"{capacity_duplicates}"
    )

    print(
        f"outage duplicate report names skipped: "
        f"{outage_duplicates}"
    )

    print()
    print(
        "BUILDING MONTHLY CSV FILES"
    )
    print()

    for month in month_range(
        start,
        latest,
    ):

        print(
            month_string(*month)
        )

        build_capacity_month(
            month,
            capacity_by_month.get(
                month,
                [],
            ),
        )

        build_outage_month(
            month,
            outage_by_month.get(
                month,
                [],
            ),
        )

    print()
    print("Done.")
    print()
    print(
        "Second-stage extracted files are preserved under:"
    )
    print(
        "  data/ercot/capacity/extracted2/"
    )
    print(
        "  data/ercot/outages/extracted2/"
    )


if __name__ == "__main__":
    main()
