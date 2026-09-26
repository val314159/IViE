#!/usr/bin/env python3
"""Download five years of ERCOT history, newest month first. Python 3.10+.

Dependency: python3 -m pip install openpyxl

Quick start (no database connection or ingestion):
    python3 download_ercot.py
    python3 download_ercot.py --months 60 --end-month 2026-08 --extract
    python3 download_ercot.py --months 13 --end-month 2026-08 --extract
    python3 download_ercot.py --datasets prices demand fuel_mix --extract
    python3 download_ercot.py --year 2025 --month 8 --out data/ercot
    python3 download_ercot.py --offline --datasets prices demand fuel_mix --year 2025 --month 8

Default: last 60 COMPLETE calendar months, ending last month in America/Chicago.
The most recent 13 months run first as the fallback baseline. Missing older data
does not discard newer successes or stop other datasets. range_summary.json
reports full-range and recent-13-month completion separately. Exit 1 still means
there are gaps in the requested range, even when the fallback baseline succeeded.
--months 13 requests ONLY that baseline. --year YYYY --month M retains the old
single-month behavior. --dry-run prints dates without downloads or credentials.
The selected outage product NP1-346-ER began in December 2022; earlier months
cannot be completed from that product and are reported as gaps, not fabricated
or silently replaced with a different report.

Prices, demand and fuel mix: cache original annual files, then export ONLY the
each selected month's rows to <out>/YYYY-MM/{prices,demand,fuel_mix}.csv. Original
column names, values, and time conventions are preserved. This is file slicing,
not database ingestion. Fuel mix retains its daily rows / quarter-hour columns.
Outages and capacity: download historical report archives for each month.
Outages include publication dates shifted +3 days, because that report describes
the third day before publication. Capacity includes 7 preceding publication days
to retain forecasts covering the start of the month. These are report snapshots,
not reconstructed actual available capacity. All versions/formats are retained.

Historical outages/capacity require a registered ERCOT Public API account:
    https://developer.ercot.com/applications/pubapi/user-guide/registration-and-authentication/
Set ERCOT_SUBSCRIPTION_KEY and either ERCOT_ID_TOKEN or both ERCOT_USERNAME and
ERCOT_PASSWORD in your environment. Do not commit credentials. With username /
password the script obtains and renews the token automatically. An externally
provided ID token expires after about an hour; rerun with a new one if necessary.

Source files are retained under data/ercot/<dataset>/raw/. --extract also unpacks ZIPs
one level under <dataset>/extracted/ (nested ZIPs remain intact). XLSX files are
kept as spreadsheets. manifest.json records URLs, hashes, report IDs, requested
period, and per-dataset status. Credentials are checked BEFORE an all-five run;
no partial run starts if archive access is unavailable. --offline works only with
an explicit selection of prices/demand/fuel_mix and reuses verified cached files.
Reruns verify hashes and skip completed downloads;
--force downloads them again. Exit 0 means all selected datasets succeeded;
exit 1 means at least one was missing, blocked or failed. No fake data is created.

Sources / protocol references (checked September 2026):
    https://www.ercot.com/mp/data-products/data-product-details?id=NP6-785-ER
    https://www.ercot.com/gridinfo/load/load_hist
    https://www.ercot.com/mp/data-products/data-product-details?id=NP1-346-ER
    https://www.ercot.com/mp/data-products/data-product-details?id=NP3-763-CD
    https://www.ercot.com/gridinfo/generation
    https://github.com/ercot/api-specs/discussions/39
"""

import argparse
import calendar
import csv
import hashlib
import json
import os
import re
import shutil
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
import zipfile
from zoneinfo import ZoneInfo
from datetime import date, datetime, timedelta, timezone
from html.parser import HTMLParser
from pathlib import Path

BASE = "https://www.ercot.com"
API = "https://api.ercot.com/api/public-reports"
TOKEN_URL = ("https://ercotb2c.b2clogin.com/ercotb2c.onmicrosoft.com/"
             "B2C_1_PUBAPI-ROPC-FLOW/oauth2/v2.0/token")
CLIENT_ID = "fec253ea-0d06-4272-a5e6-b478baeecd70"
DATASETS = ("prices", "demand", "outages", "capacity", "fuel_mix")
PAGES = {"demand": BASE + "/gridinfo/load/load_hist",
         "fuel_mix": BASE + "/gridinfo/generation"}
PRODUCTS = {"outages": "NP1-346-ER", "capacity": "NP3-763-CD"}


class DownloadError(Exception):
    pass


def utc_now():
    return datetime.now(timezone.utc).isoformat()


def sha256(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def write_json(path, value):
    temp = path.with_name(path.name + ".part")
    temp.write_text(json.dumps(value, indent=2) + "\n", encoding="utf-8")
    temp.replace(path)


class Links(HTMLParser):
    """Read both title attributes and visible anchor text, without extra packages."""
    def __init__(self):
        super().__init__()
        self.links = []
        self.current = None

    def handle_starttag(self, tag, attrs):
        if tag == "a":
            attrs = dict(attrs)
            self.current = [attrs.get("href", ""), attrs.get("title", "")]

    def handle_data(self, text):
        if self.current is not None:
            self.current[1] += " " + text

    def handle_endtag(self, tag):
        if tag == "a" and self.current is not None:
            self.links.append(tuple(self.current))
            self.current = None


class Client:
    def __init__(self, timeout=120):
        self.timeout = timeout
        self.key = os.getenv("ERCOT_SUBSCRIPTION_KEY", "")
        self.token = os.getenv("ERCOT_ID_TOKEN", "")
        self.user = os.getenv("ERCOT_USERNAME", "")
        self.password = os.getenv("ERCOT_PASSWORD", "")
        self.expires = time.monotonic() + 3300 if self.token else 0
        self.last_api_call = 0

    def authenticate(self):
        if not self.key:
            raise DownloadError("Needs ERCOT_SUBSCRIPTION_KEY plus ERCOT_ID_TOKEN "
                                "or ERCOT_USERNAME / ERCOT_PASSWORD; see script docstring.")
        if self.token and time.monotonic() < self.expires:
            return
        if not (self.user and self.password):
            raise DownloadError("Missing or expired ERCOT_ID_TOKEN; supply a fresh token "
                                "or set ERCOT_USERNAME and ERCOT_PASSWORD.")
        body = urllib.parse.urlencode({
            "username": self.user, "password": self.password,
            "grant_type": "password", "scope": f"openid {CLIENT_ID} offline_access",
            "client_id": CLIENT_ID, "response_type": "id_token",
        }).encode()
        with self.open(TOKEN_URL, body, "application/x-www-form-urlencoded") as r:
            payload = json.load(r)
        self.token = payload.get("id_token", "")
        if not self.token:
            raise DownloadError("ERCOT authentication did not return an ID token.")
        self.expires = time.monotonic() + max(1, int(payload.get("expires_in", 3600)) - 120)

    def open(self, url, body=None, content_type=None, authenticated=False):
        if authenticated and (urllib.parse.urlsplit(url).scheme != "https" or
                              urllib.parse.urlsplit(url).netloc != "api.ercot.com"):
            raise DownloadError("Refusing to send API credentials to an unexpected host.")
        for attempt in range(4):
            if authenticated:
                self.authenticate()
                time.sleep(max(0, 2.1 - (time.monotonic() - self.last_api_call)))
                self.last_api_call = time.monotonic()
            req = urllib.request.Request(url, data=body,
                                         headers={"User-Agent": "Agatha-ERCOT-Downloader/1.0"})
            if content_type:
                req.add_header("Content-Type", content_type)
            if authenticated:
                # These headers must not follow a redirect to a different host.
                req.add_unredirected_header("Authorization", "Bearer " + self.token)
                req.add_unredirected_header("Ocp-Apim-Subscription-Key", self.key)
            try:
                return urllib.request.urlopen(req, timeout=self.timeout)
            except urllib.error.HTTPError as exc:
                status = exc.code
                retry_after = exc.headers.get("Retry-After", "")
                exc.close()
                if status == 401 and authenticated and self.user and self.password and attempt < 3:
                    self.token = ""
                    continue
                if status not in (429, 500, 502, 503, 504) or attempt == 3:
                    raise DownloadError(f"HTTP {status} from {urllib.parse.urlsplit(url).hostname}") from None
                delay = min(30, int(retry_after)) if retry_after.isdigit() else 2 ** attempt
            except (urllib.error.URLError, TimeoutError, OSError):
                if attempt == 3:
                    raise DownloadError("Network request failed after four attempts.") from None
                delay = 2 ** attempt
            time.sleep(delay)

    def json(self, url, authenticated=False):
        with self.open(url, authenticated=authenticated) as response:
            return json.load(response)


def annual_links(client, dataset, year):
    page = PAGES[dataset]
    with client.open(page) as r:
        html = r.read().decode("utf-8")
    parser = Links()
    parser.feed(html)
    matches = []
    for href, label in parser.links:
        target = "hourly load data" if dataset == "demand" else "fuel mix report"
        if target not in label.lower():
            continue
        years = [int(y) for y in re.findall(r"\b(?:19|20)\d{2}\b", label)]
        if not years or not min(years) <= year <= max(years):
            continue
        url = urllib.parse.urljoin(page, href)
        if not re.search(r"\.(zip|xlsx|xls)$", urllib.parse.urlsplit(url).path, re.I):
            continue
        matches.append({"url": url, "filename": Path(urllib.parse.urlsplit(url).path).name,
                        "source_page": page, "coverage_years": sorted(set(years)),
                        "scope": "Original annual file or multi-year bundle; no row filtering."})
    if not matches:
        raise DownloadError(f"No {year} file found on {page}; source layout or availability changed.")
    # Prefer a dedicated annual file over a larger multi-year bundle.
    exact = [x for x in matches if x["coverage_years"] == [year]]
    return list({x["url"]: x for x in (exact or matches)}.values())


def price_links(client, year):
    listing = BASE + "/misapp/servlets/IceDocListJsonWS?reportTypeId=13061"
    payload = client.json(listing)
    rows = payload["ListDocsByRptTypeRes"]["DocumentList"]
    docs = [r["Document"] for r in rows
            if re.search(rf"(?<!\d){year}(?!\d)", r["Document"].get("FriendlyName", ""))]
    if not docs:
        raise DownloadError(f"No public annual RTM price file found for {year}.")
    doc = max(docs, key=lambda d: d.get("PublishDate", ""))
    return [{"url": BASE + "/misdownload/servlets/mirDownload?" + urllib.parse.urlencode({
                "doclookupId": doc["DocID"]}),
             "filename": doc["ConstructedName"], "doc_id": doc["DocID"],
             "published_at": doc["PublishDate"], "source_page": listing,
             "coverage_year": year, "scope": "Full annual price file; no row filtering."}]


def archive_window(dataset, year, month):
    start = date(year, month, 1)
    end = date(year, month, calendar.monthrange(year, month)[1])
    if dataset == "outages":
        return start + timedelta(days=3), end + timedelta(days=3)
    return start - timedelta(days=7), end


def archive_links(client, dataset, year, month):
    product = PRODUCTS[dataset]
    start, end = archive_window(dataset, year, month)
    endpoint = f"{API}/archive/{product}"
    docs, seen = [], set()
    page = 1
    while True:
        params = {"postDatetimeFrom": f"{start}T00:00:00",
                  "postDatetimeTo": f"{end}T23:59:59", "page": page, "size": 1000}
        payload = client.json(endpoint + "?" + urllib.parse.urlencode(params), authenticated=True)
        if "archives" not in payload or "totalPages" not in payload.get("_meta", {}):
            raise DownloadError("Unexpected archive response; refusing to report a complete download.")
        for doc in payload["archives"]:
            if str(doc["docId"]) not in seen:
                docs.append(doc)
                seen.add(str(doc["docId"]))
        if page >= int(payload["_meta"]["totalPages"]):
            break
        if not payload["archives"]:
            raise DownloadError("Empty intermediate archive page; results may be incomplete.")
        page += 1
    if not docs:
        raise DownloadError(f"No {product} archives returned for publication dates {start} to {end}.")
    # Small batches reduce HTTP overhead while staying below typical bulk limits.
    docs.sort(key=lambda d: (d["postDatetime"], str(d["docId"])))
    batches = []
    for offset in range(0, len(docs), 20):
        group = docs[offset:offset + 20]
        ids = [d["docId"] for d in group]
        key = hashlib.sha256(json.dumps(ids).encode()).hexdigest()[:12]
        batches.append({"url": endpoint + "/download", "body": {"docIds": ids},
                        "authenticated": True,
                        "filename": f"{product}_{year}-{month:02d}_{key}.zip",
                        "publication_window": [str(start), str(end)],
                        "documents": group, "scope": "Original archived report snapshots."})
    return batches


def validate_file(path, suffix):
    with path.open("rb") as f:
        prefix = f.read(512).lstrip().lower()
    if not prefix or prefix.startswith((b"<!doctype html", b"<html", b"{", b"[")):
        raise DownloadError("Server returned an empty file or an HTML/JSON response instead of data.")
    if suffix.lower() in (".zip", ".xlsx"):
        try:
            with zipfile.ZipFile(path) as archive:
                bad = archive.testzip()
                if bad:
                    raise DownloadError("ZIP integrity check failed.")
        except zipfile.BadZipFile:
            raise DownloadError("Expected a ZIP/XLSX file, but the response is not a valid archive.") from None


def extract_zip(path, destination):
    """Extract one level; reject traversal/symlinks and limit uncompressed size."""
    destination = destination.resolve()
    with zipfile.ZipFile(path) as archive:
        if sum(i.file_size for i in archive.infolist()) > 5 * 1024 ** 3:
            raise DownloadError("Archive exceeds the 5 GiB extraction limit.")
        intact = True
        for item in archive.infolist():
            target = (destination / item.filename).resolve()
            if not target.is_relative_to(destination) or (item.external_attr >> 16) & 0o170000 == 0o120000:
                raise DownloadError("Unsafe path in ZIP archive.")
            if not item.is_dir() and (not target.is_file() or target.stat().st_size != item.file_size):
                intact = False
        marker = destination / ".source-sha256"
        digest = sha256(path)
        if intact and marker.is_file() and marker.read_text() == digest:
            return
        destination.mkdir(parents=True, exist_ok=True)
        archive.extractall(destination)
        marker.write_text(digest)


def download(client, job, dataset, root, manifest, force=False, extract=False):
    filename = job["filename"]
    if filename != Path(filename).name or filename in ("", ".", ".."):
        raise DownloadError("Unexpected filename from source.")
    relative = str(Path(dataset) / "raw" / filename)
    target = root / relative
    target.parent.mkdir(parents=True, exist_ok=True)
    old = manifest["files"].get(relative, {})
    if not force and target.is_file() and old.get("sha256") == sha256(target):
        print(f"  SKIP {relative}", flush=True)
    else:
        print(f"  GET  {relative}", flush=True)
        temp = target.with_name(target.name + ".part")
        try:
            body = json.dumps(job["body"]).encode() if "body" in job else None
            with client.open(job["url"], body, "application/json" if body else None,
                             job.get("authenticated", False)) as response, temp.open("wb") as f:
                shutil.copyfileobj(response, f, length=1024 * 1024)
            validate_file(temp, target.suffix)
            temp.replace(target)
        finally:
            temp.unlink(missing_ok=True)
        manifest["files"][relative] = {**{k: v for k, v in job.items() if k != "authenticated"},
                                       "downloaded_at": utc_now(), "bytes": target.stat().st_size,
                                       "sha256": sha256(target)}
        write_json(root / "manifest.json", manifest)
    if extract and target.suffix.lower() == ".zip":
        extract_zip(target, root / dataset / "extracted" / target.stem)
    return target


def cached_sources(root, manifest, dataset, year):
    """Use only source files with matching provenance and verified checksums."""
    paths = []
    for relative, meta in manifest["files"].items():
        if not relative.startswith(dataset + "/raw/"):
            continue
        years = meta.get("coverage_years", [meta.get("coverage_year")])
        years = [y for y in years if isinstance(y, int)]
        if not years or not min(years) <= year <= max(years):
            continue
        path = root / relative
        if not path.is_file() or sha256(path) != meta.get("sha256"):
            raise DownloadError(f"Missing or modified cached source: {relative}; rerun without --offline.")
        paths.append(path)
    if len(paths) != 1:
        raise DownloadError(f"Expected one verified {year} source for {dataset}; found {len(paths)}. "
                            "Rerun without --offline to resolve the current source.")
    return paths


def source_date(value):
    # Use the operating date on the source row. In particular, keep the day's
    # 24:00 hour-ending observation with THAT day, not with the following day.
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    value = str(value).strip()
    for pattern, fmt in ((r"^\d{2}/\d{2}/\d{4}", "%m/%d/%Y"),
                         (r"^\d{4}-\d{2}-\d{2}", "%Y-%m-%d")):
        match = re.match(pattern, value)
        if match:
            return datetime.strptime(match.group(), fmt).date()
    raise DownloadError(f"Unrecognized date in source: {value!r}")


def monthly_csv(paths, dataset, year, month, root):
    """Preserve the original table layout while selecting the requested month."""
    from openpyxl import load_workbook
    output_dir = root / f"{year}-{month:02d}"
    output_dir.mkdir(parents=True, exist_ok=True)
    output = output_dir / f"{dataset}.csv"
    temp = output.with_suffix(".csv.part")
    expected_header = {"prices": "Delivery Date", "demand": "Hour Ending", "fuel_mix": "Date"}[dataset]
    count, days, source_files, header = 0, set(), [], None
    try:
        with temp.open("w", newline="", encoding="utf-8") as target:
            writer = csv.writer(target)
            for source in paths:
                if source.suffix.lower() == ".zip":
                    folder = root / dataset / "extracted" / source.stem
                    extract_zip(source, folder)
                    workbooks = sorted(folder.rglob("*.xlsx"))
                else:
                    workbooks = [source]
                if not workbooks:
                    raise DownloadError("No XLSX workbook found in the downloaded source archive.")
                for path in workbooks:
                    if path.suffix.lower() != ".xlsx":
                        raise DownloadError("Monthly export currently supports XLSX sources only.")
                    # Historical bundles can contain multiple years. Only inspect
                    # files naming our year; dedicated annual filenames also do so.
                    named_years = re.findall(r"(?<!\d)(?:19|20)\d{2}(?!\d)", path.stem)
                    if named_years and str(year) not in named_years:
                        continue
                    wb = load_workbook(path, read_only=True, data_only=True)
                    try:
                        if dataset in ("prices", "fuel_mix"):
                            name = calendar.month_abbr[month]
                            if name not in wb.sheetnames:
                                raise DownloadError(f"Expected month sheet {name!r} in {path.name}.")
                            sheets = [wb[name]]
                        else:
                            sheets = list(wb)
                        for sheet in sheets:
                            rows = sheet.iter_rows(values_only=True)
                            sheet_header = next(rows, ())
                            if not sheet_header or sheet_header[0] != expected_header:
                                raise DownloadError(f"Unexpected header in {path.name}/{sheet.title}.")
                            if header is None:
                                header = sheet_header
                                writer.writerow(header)
                            elif header != sheet_header:
                                raise DownloadError("Source headers differ; refusing to merge incompatible tables.")
                            source_files.append(str(path.relative_to(root)))
                            for row in rows:
                                if not row or all(v is None for v in row):
                                    continue
                                day = source_date(row[0])
                                if (day.year, day.month) == (year, month):
                                    writer.writerow([v.isoformat() if isinstance(v, (date, datetime)) else v
                                                     for v in row])
                                    days.add(day)
                                    count += 1
                    finally:
                        wb.close()
        expected_days = {date(year, month, n) for n in range(1, calendar.monthrange(year, month)[1] + 1)}
        if days != expected_days:
            missing = sorted(str(day) for day in expected_days - days)
            raise DownloadError(f"Monthly file does not cover every date; missing: {', '.join(missing)}")
        temp.replace(output)
    finally:
        temp.unlink(missing_ok=True)
    print(f"  MONTH {output.relative_to(root)}: {count:,} rows, {min(days)} through {max(days)}", flush=True)
    return {"path": str(output.relative_to(root)), "rows": count,
            "first_date": str(min(days)), "last_date": str(max(days)),
            "source_files": source_files, "sha256": sha256(output),
            "time_convention": "Source operating dates and hour-ending labels retained, including 24:00."}


def month_range(count, end_month=None, today=None):
    """Newest first, including exactly count complete calendar months."""
    today = today or datetime.now(ZoneInfo("America/Chicago")).date()
    last_complete = today.replace(day=1) - timedelta(days=1)
    if end_month:
        if not re.fullmatch(r"\d{4}-\d{2}", end_month):
            raise ValueError("--end-month must be YYYY-MM")
        end = date.fromisoformat(end_month + "-01")
    else:
        end = last_complete.replace(day=1)
    if end > last_complete or count < 1 or count > 120:
        raise ValueError("Choose 1–120 months ending no later than the last complete month.")
    index = end.year * 12 + end.month - 1
    periods = [divmod(index - n, 12) for n in range(count)]
    periods = [(year, month + 1) for year, month in periods]
    if periods[-1][0] < 1995:
        raise ValueError("Requested history starts before 1995.")
    return periods


def range_summary(run, periods, datasets):
    def coverage(selected):
        gaps = []
        for year, month in selected:
            key = f"{year}-{month:02d}"
            results = run["periods"].get(key, {})
            for dataset in datasets:
                item = results.get(dataset, {})
                if item.get("status") != "complete":
                    gaps.append({"month": key, "dataset": dataset,
                                 "reason": item.get("error", item.get("status", "not attempted"))})
        return {"start_month": "%04d-%02d" % selected[-1],
                "end_month": "%04d-%02d" % selected[0], "months": len(selected),
                "complete": not gaps, "gaps": gaps}
    return {"updated_at": utc_now(), "selected_datasets": datasets,
            "requested_range": coverage(periods), "recent_baseline": coverage(periods[:13]),
            "note": "Public datasets are month-only CSVs. Archives are source report snapshots; "
                    "archive checks verify listed files and publication-day coverage, not every intraday observation."}


def archive_missing_days(jobs, dataset, year, month):
    start, end = archive_window(dataset, year, month)
    seen = {date.fromisoformat(doc["postDatetime"][:10])
            for job in jobs for doc in job["documents"]}
    return [str(start + timedelta(days=n)) for n in range((end - start).days + 1)
            if start + timedelta(days=n) not in seen]


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--months", type=int, help="Number of complete months (default: 60).")
    parser.add_argument("--end-month", help="Last month, YYYY-MM (default: previous calendar month).")
    parser.add_argument("--year", type=int, help="Legacy single-month mode; use with --month.")
    parser.add_argument("--month", type=int, choices=range(1, 13), help="Legacy single-month mode; use with --year.")
    parser.add_argument("--datasets", nargs="+", choices=DATASETS, default=list(DATASETS))
    parser.add_argument("--out", type=Path, default=Path("data/ercot"))
    parser.add_argument("--extract", action="store_true", help="Also unpack downloaded archive ZIPs one level.")
    parser.add_argument("--force", action="store_true", help="Redownload completed files and rebuild monthly CSVs.")
    parser.add_argument("--offline", action="store_true", help="Export monthly CSVs from verified cached public sources.")
    parser.add_argument("--dry-run", action="store_true", help="Print range and ordering without network or file writes.")
    parser.add_argument("--timeout", type=int, default=120)
    args = parser.parse_args(argv)
    if args.timeout <= 0:
        parser.error("Use a positive timeout.")
    try:
        if args.year is not None or args.month is not None:
            if args.year is None or args.month is None or args.months is not None or args.end_month is not None:
                parser.error("Use --year AND --month for one month, OR --months / --end-month for a range.")
            periods = month_range(1, f"{args.year:04d}-{args.month:02d}")
        else:
            periods = month_range(args.months if args.months is not None else 60, args.end_month)
    except ValueError as exc:
        parser.error(str(exc))
    datasets = list(dict.fromkeys(args.datasets))
    first, last = "%04d-%02d" % periods[-1], "%04d-%02d" % periods[0]
    baseline_start = "%04d-%02d" % periods[min(12, len(periods) - 1)]
    print(f"Requested: {first} through {last} ({len(periods)} months), newest first.", flush=True)
    print(f"Priority baseline: {baseline_start} through {last} ({min(13, len(periods))} months).", flush=True)
    if args.dry_run:
        print("Datasets: " + ", ".join(datasets))
        print("Order: " + ", ".join("%04d-%02d" % p for p in periods))
        return 0
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        parser.error("Monthly file export needs openpyxl: python3 -m pip install openpyxl")
    client = Client(args.timeout)
    archives = [d for d in datasets if d in PRODUCTS]
    if args.offline and archives:
        parser.error("--offline supports --datasets prices demand fuel_mix; archives need API access.")
    if archives:
        try:
            client.authenticate()
            for dataset in archives:
                # Verify authorization, independent of the requested dates. A
                # missing older month is recorded later, not a reason to abort.
                payload = client.json(f"{API}/archive/{PRODUCTS[dataset]}?page=1&size=1", authenticated=True)
                if "archives" not in payload:
                    raise DownloadError(f"Unexpected {dataset} archive response during access check.")
        except (DownloadError, OSError, ValueError) as exc:
            print(f"PRECHECK FAILED: {exc}\nNo downloads started. For public sources only, explicitly use "
                  "--datasets prices demand fuel_mix.", file=sys.stderr)
            return 1
    root = args.out.resolve()
    root.mkdir(parents=True, exist_ok=True)
    path = root / "manifest.json"
    manifest = json.loads(path.read_text()) if path.exists() else {"files": {}, "runs": []}
    monthly_history = manifest.setdefault("monthly_exports", {})
    run = {"started_at": utc_now(), "start_month": first, "end_month": last, "periods": {}}
    manifest["runs"].append(run)
    annual_cache, fetched_jobs = {}, {}

    def checkpoint():
        write_json(path, manifest)
        summary = range_summary(run, periods, datasets)
        write_json(root / "range_summary.json", summary)
        return summary

    checkpoint()
    for index, (year, month) in enumerate(periods, 1):
        period = f"{year}-{month:02d}"
        print(f"\n=== {period} ({index}/{len(periods)}) ===", flush=True)
        results = run["periods"][period] = {}
        for dataset in datasets:
            print(f"\n{dataset.upper()}", flush=True)
            result = results[dataset] = {"status": "running", "files_completed": 0}
            try:
                if dataset not in PRODUCTS:
                    cache_key = (dataset, year)
                    if cache_key not in annual_cache:
                        if args.offline:
                            sources = cached_sources(root, manifest, dataset, year)
                        else:
                            jobs = price_links(client, year) if dataset == "prices" else annual_links(client, dataset, year)
                            sources = []
                            for job in jobs:
                                # A multi-year fuel archive may serve several years.
                                identity = (dataset, job["url"])
                                if identity not in fetched_jobs:
                                    fetched_jobs[identity] = download(client, job, dataset, root, manifest,
                                                                       args.force, args.extract)
                                sources.append(fetched_jobs[identity])
                        annual_cache[cache_key] = sources
                    sources = annual_cache[cache_key]
                    fingerprints = {str(p.relative_to(root)): manifest["files"][str(p.relative_to(root))]["sha256"]
                                    for p in sources}
                    export_key = f"{period}/{dataset}"
                    previous = monthly_history.get(export_key, {})
                    monthly = previous.get("monthly_file", {})
                    output = root / monthly.get("path", "__not_present__")
                    if (not args.force and previous.get("inputs") == fingerprints and output.is_file()
                            and monthly.get("sha256") == sha256(output)):
                        print(f"  SKIP MONTH {monthly['path']}", flush=True)
                        result["monthly_file"] = monthly
                    else:
                        result["monthly_file"] = monthly_csv(sources, dataset, year, month, root)
                        monthly_history[export_key] = {"inputs": fingerprints, "monthly_file": result["monthly_file"]}
                    result["files_completed"] = len(sources)
                else:
                    if dataset == "outages" and archive_window(dataset, year, month)[1] < date(2022, 12, 8):
                        raise DownloadError("NP1-346-ER starts in December 2022; this month is unavailable from that report.")
                    jobs = archive_links(client, dataset, year, month)
                    result["files_expected"] = len(jobs)
                    result["archive_files"] = []
                    for batch_number, job in enumerate(jobs, 1):
                        print(f"  Bundle {batch_number}/{len(jobs)}", flush=True)
                        output = download(client, job, dataset, root, manifest, args.force, args.extract)
                        result["archive_files"].append(str(output.relative_to(root)))
                        result["files_completed"] += 1
                    result["scope"] = "Archived snapshots, with publication-date padding; not consolidated monthly rows."
                    missing = archive_missing_days(jobs, dataset, year, month)
                    if missing:
                        result["missing_publication_dates"] = missing
                        raise DownloadError("Archives missing publication dates: " + ", ".join(missing))
                result["status"] = "complete"
            except (DownloadError, OSError, ValueError, KeyError, zipfile.BadZipFile) as exc:
                result.update(status="incomplete", error=str(exc))
                print(f"  INCOMPLETE: {exc}", file=sys.stderr, flush=True)
            checkpoint()
        if index == min(13, len(periods)):
            summary = range_summary(run, periods, datasets)
            print("\nPRIORITY BASELINE: " + ("COMPLETE" if summary["recent_baseline"]["complete"] else
                                           "INCOMPLETE — see range_summary.json"), flush=True)
    run["finished_at"] = utc_now()
    summary = checkpoint()
    print(f"\nRange report: {root / 'range_summary.json'}")
    for label, key in (("Requested range", "requested_range"), ("Recent baseline", "recent_baseline")):
        section = summary[key]
        print(f"{label}: {section['start_month']} through {section['end_month']} — "
              f"{'COMPLETE' if section['complete'] else 'INCOMPLETE'} ({len(section['gaps'])} dataset/month gaps)")
    return 0 if summary["requested_range"]["complete"] else 1

if __name__ == "__main__":
    try:
        sys.exit(main())
    except KeyboardInterrupt:
        print("\nInterrupted. Rerun to skip verified completed files.", file=sys.stderr)
        sys.exit(130)
